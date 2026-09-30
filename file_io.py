import csv
from pathlib import Path

from openpyxl import Workbook

#Чтение CSV, проверка заголовка и возврат строк 
def read_csv(path: Path) -> tuple[list[str], list[tuple[int, list[str]]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file, delimiter=",")
        headers = next(reader, None)
        if headers is None:
            raise ValueError("CSV пуст")
        if len(headers) != 9:
            raise ValueError(f"В заголовке {len(headers)} столбцов вместо 9: {headers}")
        rows = list(enumerate(reader, start=2))
    return headers, rows

#Добавление строки на лист
def add_excel_row(sheet, values: list) -> None:
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"

#Создание новых листов  и запись их
def write_excel(
    path: Path, headers: list[str], clean_rows: list[list], error_rows: list[list]
) -> None:
    """Создаёт листы, форматирует даты и суммы и сохраняет книгу."""
    workbook = Workbook()
    clean_sheet = workbook.active
    clean_sheet.title = "Корректные"
    error_sheet = workbook.create_sheet("Ошибки")
    add_excel_row(clean_sheet, headers)
    add_excel_row(error_sheet, headers + ["Тип ошибки", "Строка CSV", "Лишние поля"])

    for cleaned in clean_rows:
        add_excel_row(clean_sheet, cleaned)
        clean_sheet.cell(clean_sheet.max_row, 2).number_format = "yyyy-mm-dd"
        clean_sheet.cell(clean_sheet.max_row, 6).number_format = "#,##0.00"
    for problem in error_rows:
        add_excel_row(error_sheet, problem)

    workbook.save(path)
